"""Retain and extract the predeclared CAN/USA 1950-1970 annual observations.

No network, workbook authoring, fitting, interpolation, or holdout search.
Run --check to verify the portable retained inputs and exact derived JSON.
"""
from pathlib import Path
from hashlib import sha256
from collections import Counter
import argparse
import json
import math
import openpyxl

ROOT = Path(__file__).resolve().parent
ORIGINAL = Path('B:/Projects/SOT_LAKE/LAKE/out')
RETAINED = ROOT / 'data/local-signal'
WORKBOOK_NAME = 'mpd2023_web.xlsx'
MANIFEST_NAME = 'manifest_pull.json'
EXPECTED_WORKBOOK_SHA256 = 'ecc5916ca12789b983fc4be437f8a354bbf4291323605324ac3e0aea4c57cbb6'
EXPECTED_MANIFEST_SHA256 = 'a9be4d88cb4fff60dea6885f341295636db325cf0ab77253bd65be65e9b27467'
COUNTRIES = ('CAN', 'USA')
START_YEAR, END_YEAR = 1950, 1970


def require(condition, message):
    if not condition:
        raise ValueError(message)


def digest(path):
    return sha256(Path(path).read_bytes()).hexdigest()


def json_bytes(value):
    return (json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + '\n').encode('utf-8')


def retain_input(original, retained, expected):
    """Idempotent byte preservation; a different existing copy is never overwritten."""
    if retained.exists():
        require(digest(retained) == expected, 'Retained source revision disagrees: ' + str(retained))
        return
    raw = original.read_bytes()
    require(sha256(raw).hexdigest() == expected, 'Original source revision disagrees: ' + str(original))
    retained.parent.mkdir(parents=True, exist_ok=True)
    with retained.open('xb') as output:
        output.write(raw)
    require(digest(retained) == expected, 'Retained source bytes changed during copy.')


def metadata_rows(sheet, include_all=False):
    selected, active = [], None
    for row_number, row in enumerate(sheet.iter_rows(max_col=sheet.max_column, values_only=False), 1):
        first = row[0].value
        if first is not None:
            active = first
        if include_all or active in COUNTRIES or (sheet.title == 'Sources' and row_number <= 2):
            cells = [{'cell': cell.coordinate, 'value': cell.value, 'excelType': cell.data_type}
                     for cell in row if cell.value is not None]
            if cells:
                selected.append({'sheet': sheet.title, 'row': row_number, 'cells': cells})
    return selected


def read_wide_sheet(sheet):
    # Only the two declared country columns enter the extracted record scope.
    headers = list(sheet.iter_rows(min_row=1, max_row=3, values_only=True))
    result = {}
    for code in COUNTRIES:
        positions = [index for index, value in enumerate(headers[2]) if value == code]
        require(len(positions) == 1, 'Country identity is absent or duplicated: ' + code)
        column = positions[0]
        result[code] = {'name': headers[0][column], 'region': headers[1][column],
                        'column': openpyxl.utils.get_column_letter(column + 1), 'records': []}
    for cells in sheet.iter_rows(min_row=4, max_row=sheet.max_row):
        year = cells[0].value
        if not isinstance(year, (int, float)) or isinstance(year, bool) or not START_YEAR <= year <= END_YEAR:
            continue
        require(float(year).is_integer(), 'Fractional source year is unsupported.')
        for code in COUNTRIES:
            column = openpyxl.utils.column_index_from_string(result[code]['column']) - 1
            cell = cells[column]
            value = cell.value
            require(value is None or (isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)),
                    'Unexpected nonnumeric source value at ' + sheet.title + '!' + cell.coordinate)
            result[code]['records'].append({'year': int(year), 'value': value,
                                            'yearCell': cells[0].coordinate, 'valueCell': cell.coordinate,
                                            'excelType': cell.data_type})
    return result


def sampling_record(records):
    years = [record['year'] for record in records]
    counts = Counter(years)
    steps = [b-a for a, b in zip(years, years[1:])]
    expected = list(range(START_YEAR, END_YEAR + 1))
    missing = {key: [record['year'] for record in records if record['missingness'][key]]
               for key in ('gdppc', 'pop')}
    return {'recordCount': len(records), 'timeField': 'year', 'timeUnit': 'calendar year',
            'firstYear': years[0], 'lastYear': years[-1], 'strictlyIncreasing': all(step > 0 for step in steps),
            'duplicateYears': [year for year, count in counts.items() if count > 1],
            'observedYearSteps': sorted(set(steps)), 'nominalStepYears': 1,
            'missingExpectedYears': [year for year in expected if year not in counts],
            'completeDeclaredYearGrid': years == expected, 'missingValuesByChannel': missing,
            'nonfiniteValueCount': 0, 'sourceCellFormulas': 0, 'interpolatedValues': 0,
            'synthesizedDates': 0, 'resampling': 'none', 'frequencyHz': None,
            'timeSemantics': 'Annual GDP per capita estimate; population is mid-year. A year label is not a shared instantaneous timestamp.',
            'validationScope': 'Only the declared two-country 1950-1970 window; no assertion about the full historical series.'}


def extract(workbook=RETAINED / WORKBOOK_NAME, manifest_path=RETAINED / MANIFEST_NAME):
    require(digest(workbook) == EXPECTED_WORKBOOK_SHA256, 'Unexpected workbook source revision.')
    require(digest(manifest_path) == EXPECTED_MANIFEST_SHA256, 'Unexpected pull manifest revision.')
    manifest = json.loads(Path(manifest_path).read_text(encoding='utf-8-sig'))
    events = [event for event in manifest['events'] if event.get('source') == 'maddison_dataverse_nl'
              and event.get('pile') == WORKBOOK_NAME]
    require(len(events) == 1, 'Workbook retrieval event is ambiguous or missing.')
    book = openpyxl.load_workbook(workbook, read_only=True, data_only=False)
    try:
        require(book['Notes']['A1'].value == 'Maddison Project Database (MPD) 2023', 'Unexpected source title.')
        units = {'year': 'calendar year', 'gdppc': book['Notes']['B16'].value, 'pop': book['Notes']['B17'].value}
        require(units['gdppc'] == 'Real GDP per capita in 2011$' and units['pop'] == 'Population, mid-year (thousands)', 'Unexpected native unit definitions.')
        gdp, population = read_wide_sheet(book['GDPpc']), read_wide_sheet(book['Population'])
        countries = []
        for code in COUNTRIES:
            a, b = gdp[code], population[code]
            require((a['name'], a['region'], a['column']) == (b['name'], b['region'], b['column']), 'Channel country identities disagree.')
            require([r['year'] for r in a['records']] == [r['year'] for r in b['records']], 'Native channel year grids disagree.')
            column, records = a['column'], []
            for x, y in zip(a['records'], b['records']):
                records.append({'id': 'mpd2023:' + code + ':' + str(x['year']), 'year': x['year'],
                                'values': {'gdppc': x['value'], 'pop': y['value']},
                                'sourceCells': {'year': 'GDPpc!' + x['yearCell'], 'populationYear': 'Population!' + y['yearCell'],
                                                'gdppc': 'GDPpc!' + x['valueCell'], 'pop': 'Population!' + y['valueCell'],
                                                'countrycode': 'GDPpc!' + column + '3', 'country': 'GDPpc!' + column + '1',
                                                'region': 'GDPpc!' + column + '2'},
                                'missingness': {'gdppc': x['value'] is None, 'pop': y['value'] is None},
                                'sourceTypes': {'gdppc': x['excelType'], 'pop': y['excelType']},
                                'evidenceKind': 'published historical estimates; not generated values or direct sensor samples'})
            sampling = sampling_record(records)
            require(sampling['completeDeclaredYearGrid'], 'The predeclared annual window is incomplete or unordered.')
            countries.append({'code': code, 'name': a['name'], 'region': a['region'], 'records': records, 'sampling': sampling,
                              'fieldUse': {'code': 'used: country identity', 'name': 'used: readable identity',
                                           'region': 'not used: retained native classification'}})
        notes = metadata_rows(book['Notes'], True)
        source_notes = metadata_rows(book['Sources']) + metadata_rows(book['Maddison original sources'])
        sheets = [{'name': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column} for sheet in book]
    finally:
        book.close()
    source = {'title': 'Maddison Project Database (MPD) 2023', 'revision': 'MPD2023 / doi:10.34894/INZBF2 / file 421302',
              'sha256': EXPECTED_WORKBOOK_SHA256, 'bytes': Path(workbook).stat().st_size,
              'retainedPath': 'data/local-signal/' + WORKBOOK_NAME, 'manifestPath': 'data/local-signal/' + MANIFEST_NAME,
              'manifestSha256': EXPECTED_MANIFEST_SHA256,
              'originalPath': (ORIGINAL / 'maddison_dataverse_nl' / WORKBOOK_NAME).as_posix(),
              'originalManifestPath': (ORIGINAL / MANIFEST_NAME).as_posix(),
              'retrievalEvent': events[0], 'pullStarted': manifest['started'], 'pullFinished': manifest['finished'],
              'unitCells': {'gdppc': 'Notes!B16', 'pop': 'Notes!B17'}, 'sheets': sheets,
              'notes': notes, 'countrySourceNotes': source_notes,
              'provenanceLimit': 'Country-source notes only partially resolve 1950-1970; no invented country-year citation or uncertainty estimate.',
              'duplicateFormats': 'GDPpc, Population and Full data are representations in this workbook. The accompanying Stata file is not an independent comparator.'}
    return {'version': 'local-annual-signal/1', 'id': 'mpd2023-can-usa-1950-1970',
            'title': 'Canada and United States · annual historical estimates', 'source': source,
            'window': {'startYear': START_YEAR, 'endYear': END_YEAR, 'inclusive': True,
                       'selectionRule': 'Two specified countries and a fixed 21-year window; no outcome-based search or holdout reuse for model selection.'},
            'units': units, 'countries': countries,
            'comparison': {'referenceCountry': 'CAN', 'candidateCountry': 'USA', 'alignment': 'exact native calendar-year identity only',
                           'independence': 'Distinct country observations from the same compilation; not independent data collection methods.',
                           'claim': 'Descriptive native-value comparison only; no inferred intervention, causal mechanism, phase or whole-system equivalence.'},
            'samplingValidation': {'passed': True, 'scope': '42 source rows and 84 selected values',
                                   'countryCount': 2, 'recordsPerCountry': 21, 'pairedYears': list(range(START_YEAR, END_YEAR + 1)),
                                   'sameNativeYearGrid': True, 'missingSelectedValues': sum(sum(len(v) for v in c['sampling']['missingValuesByChannel'].values()) for c in countries),
                                   'requiresSecondsTimestamp': False, 'frequencyAnalysisSupported': False, 'interpolationPermitted': False},
            'preservation': {'declaredInputScope': 'Country code/name/region, year, GDP per capita and population for CAN and USA, 1950-1970; source unit definitions and country citation metadata.',
                             'fieldCoverage': {'countrycode': 'used; retained as countries.code and record sourceCells',
                                               'country': 'used; retained as countries.name and record sourceCells',
                                               'region': 'not used; retained as countries.region and record sourceCells',
                                               'year': 'used; exact native integer and both channel year-cell pointers',
                                               'gdppc': 'used; unchanged value, missingness, unit and cell pointer',
                                               'pop': 'used; unchanged value, missingness, unit and cell pointer'},
                             'unusedOutsideScope': 'All other countries, years, regional estimates, formulas, citations, workbook formatting and metadata remain accessible in the byte-identical full workbook.',
                             'retainedFiles': [{'path': source['retainedPath'], 'sha256': EXPECTED_WORKBOOK_SHA256, 'bytes': Path(workbook).stat().st_size},
                                               {'path': source['manifestPath'], 'sha256': EXPECTED_MANIFEST_SHA256, 'bytes': Path(manifest_path).stat().st_size}],
                             'sourceModified': False, 'conversion': 'None. JSON retains native numeric values; ISO identities and integer years remain explicit.'},
            'limitations': ['Published historical estimates are not direct sensor measurements.',
                            'Annual GDP and mid-year population have different within-year meanings; a selected year does not establish simultaneous measurements.',
                            'The inspected window has no blanks; the full historical dataset contains missing and irregularly spaced observations outside this declared scope.',
                            'No per-value uncertainty, interpolation flag, or exact subannual timestamp is supplied in these selected cells.',
                            'No continuous waveform, frequency spectrum, synchronized-signal phase, audio playback, or inferred lifecycle is supported by this extraction.',
                            'Region and other unused evidence remain accessible; source notes do not establish causal relationships.']}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check', action='store_true', help='Verify retained sources and derived JSON without writing.')
    args = parser.parse_args()
    workbook, manifest = RETAINED / WORKBOOK_NAME, RETAINED / MANIFEST_NAME
    if not args.check:
        retain_input(ORIGINAL / 'maddison_dataverse_nl' / WORKBOOK_NAME, workbook, EXPECTED_WORKBOOK_SHA256)
        retain_input(ORIGINAL / MANIFEST_NAME, manifest, EXPECTED_MANIFEST_SHA256)
    result = extract(workbook, manifest)
    raw, output = json_bytes(result), ROOT / 'data/local_signal.json'
    if args.check:
        require(output.read_bytes() == raw, 'Derived local signal JSON differs from exact retained native-cell extraction.')
    else:
        output.write_bytes(raw)
    print(json.dumps({'passed': True, 'mode': 'check' if args.check else 'prepare', 'output': str(output),
                      'sha256': sha256(raw).hexdigest(), 'records': 42, 'values': 84,
                      'source_sha256': EXPECTED_WORKBOOK_SHA256, 'no_network': True, 'source_bytes_unchanged': True}, indent=2))


if __name__ == '__main__':
    main()
